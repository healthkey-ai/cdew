import base64
import io

from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import RefreshToken

from apps.accounts.models import Identity
from apps.documents.ingest import sniff_mime
from apps.documents.models import DocumentStatus, SourceDocument

# 1x1 transparent PNG (valid PNG magic signature).
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAAC0lEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Sodium", 140
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class DocumentUploadTests(APITestCase):
    def setUp(self):
        self.identity = Identity.objects.create(issuer="urn:local", sub="u1", email="cdm@x.org")
        token = RefreshToken.for_user(self.identity).access_token
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")

    def _upload(self, name, content, ctype, patient="P1"):
        f = SimpleUploadedFile(name, content, content_type=ctype)
        return self.client.post("/api/v1/documents/", {"file": f, "patient_ref": patient}, format="multipart")

    def test_upload_requires_auth(self):
        self.client.credentials()  # drop token
        r = self._upload("a.txt", b"hello world text", "text/plain")
        assert r.status_code == 401

    def test_txt_upload_parses_text(self):
        r = self._upload("labs.txt", b"Hemoglobin 9.1 g/dL (Low)\nPatient stable.", "text/plain")
        assert r.status_code == 201, r.content
        assert r.json()["has_text"] is True  # status advances past parsed once classified
        doc_id = r.json()["id"]
        t = self.client.get(f"/api/v1/documents/{doc_id}/text/")
        assert t.status_code == 200
        assert "Hemoglobin" in t.json()["content"]
        assert t.json()["pages"][0]["char_end"] == t.json()["char_len"]

    def test_xlsx_upload_parses_cells(self):
        r = self._upload(
            "sheet.xlsx", _xlsx_bytes(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        assert r.status_code == 201, r.content
        doc_id = r.json()["id"]
        content = self.client.get(f"/api/v1/documents/{doc_id}/text/").json()["content"]
        assert "Sodium" in content

    def test_dedup_same_bytes_same_patient(self):
        payload = b"Sodium 140 mmol/L reference range."
        r1 = self._upload("a.txt", payload, "text/plain", patient="P9")
        r2 = self._upload("a-again.txt", payload, "text/plain", patient="P9")
        assert r1.status_code == 201 and r1.json()["deduplicated"] is False
        assert r2.status_code == 200 and r2.json()["deduplicated"] is True
        assert r1.json()["id"] == r2.json()["id"]
        assert SourceDocument.objects.filter(patient_ref="P9").count() == 1

    def test_same_bytes_different_patient_not_deduped(self):
        payload = b"Calcium 9.5 mg/dL within range."
        self._upload("a.txt", payload, "text/plain", patient="A")
        r2 = self._upload("a.txt", payload, "text/plain", patient="B")
        assert r2.status_code == 201
        assert SourceDocument.objects.filter(sha256=r2.json()["sha256"]).count() == 2

    def test_mime_sniff_beats_extension(self):
        # PNG bytes named .pdf must be detected as image/png, not trusted as pdf.
        r = self._upload("evil.pdf", _PNG, "application/pdf")
        assert r.status_code == 201, r.content
        assert r.json()["mime"] == "image/png"

    def test_image_with_no_text_flags_for_classification(self):
        r = self._upload("scan.png", _PNG, "image/png")
        assert r.status_code == 201
        assert r.json()["status"] == DocumentStatus.NEEDS_CLASSIFICATION

    def test_unsupported_type_rejected(self):
        r = self._upload("x.bin", b"\x00\x01\x02not a real document\x03", "application/octet-stream")
        assert r.status_code == 400

    def test_phi_text_not_in_logs(self):
        with self.assertLogs("apps.documents", level="INFO") as cm:
            self._upload("labs.txt", b"Hemoglobin 9.1 g/dL SECRETVALUE42", "text/plain")
        joined = "\n".join(cm.output)
        assert "SECRETVALUE42" not in joined


class SniffMimeUnitTests(APITestCase):
    def test_png_detected_from_bytes(self):
        assert sniff_mime(_PNG, "whatever.pdf") == "image/png"

    def test_markdown_by_extension(self):
        assert sniff_mime(b"# Title\n\nbody", "notes.md") == "text/markdown"

    def test_xlsx_disambiguated_from_zip(self):
        assert sniff_mime(_xlsx_bytes(), "s.xlsx").endswith("spreadsheetml.sheet")
